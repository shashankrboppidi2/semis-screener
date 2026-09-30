"""market_bookings_screen.xlsx - implied bookings and backlog across every US filer that discloses RPO."""
import pandas as pd, numpy as np, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
import tier1_screen as TS

B = Font(bold=True); HD = PatternFill("solid", fgColor="DDE6F0"); WRAP = Alignment(wrap_text=True, vertical="top")
PCTC = {"rpo_yoy", "rpo_yoy_accel", "rev_yoy"}
MONEY = {"rev", "rpo", "d_rpo", "dr_cur", "dr_tot", "fy_rev"}

def sheet(wb, name, df, note=None, widths=None, heat=None, freeze_col=3):
    ws = wb.create_sheet(name[:31]); r = 1
    if note:
        c = ws.cell(1, 1, note); c.font = Font(italic=True); c.alignment = WRAP
        ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=max(8, len(df.columns)))
        r = 5
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(r, j, str(col)); c.font = B; c.fill = HD; c.alignment = WRAP
    for i, row in enumerate(df.itertuples(index=False), r + 1):
        for j, v in enumerate(row, 1):
            nm = str(df.columns[j - 1])
            if isinstance(v, pd.Period): v = str(v)
            if isinstance(v, (np.bool_, bool)): v = bool(v)
            c = ws.cell(i, j, None if (isinstance(v, float) and pd.isna(v)) else v)
            if isinstance(v, (int, float)) and not isinstance(v, bool) and not pd.isna(v):
                c.number_format = "0.0%" if (nm in PCTC or nm.startswith("y/y")) else ("#,##0" if nm in MONEY else "0.00")
    n = len(df)
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(str(col), max(10, min(40, len(str(col)) + 3)))
    for col in (heat or []):
        if col not in df.columns or not n: continue
        L = get_column_letter(list(df.columns).index(col) + 1)
        ws.conditional_formatting.add(f"{L}{r+1}:{L}{r+n}", ColorScaleRule(
            start_type="percentile", start_value=5, start_color="F8B3AE", mid_type="num", mid_value=1 if col.startswith("b2b") and col != "b2b_trend" else 0,
            mid_color="FFFFFF", end_type="percentile", end_value=95, end_color="A8D5A2"))
    ws.freeze_panes = ws.cell(r + 1, freeze_col)
    return ws

def traj(P, ciks, nq=9):
    rows = []
    for cik in ciks:
        g = P[(P.cik == cik)].dropna(subset=["rpo"]).sort_values("cqp").tail(nq)
        if len(g) < 4: continue
        d = {"ticker": g.ticker.iloc[-1] if "ticker" in g else None, "name": g.name.iloc[-1]}
        for _, x in g.iterrows(): d[f"RPO {x.cqp}"] = x.rpo
        for _, x in g.iterrows(): d[f"b2b {x.cqp}"] = x.b2b
        rows.append(d)
    D = pd.DataFrame(rows)
    rc = sorted([c for c in D.columns if c.startswith("RPO ")], key=lambda c: c[4:])
    bc = sorted([c for c in D.columns if c.startswith("b2b ")], key=lambda c: c[4:])
    return D[["ticker", "name"] + rc + bc]

def main():
    P, L = TS.latest()
    T = pd.read_pickle("out/cik_tickers.pkl")[["cik", "ticker"]]
    P = P.merge(T, on="cik", how="left")
    S = TS.screens(L)
    G = L[L.gated].copy()
    wb = Workbook(); wb.remove(wb.active)

    rows = [("What this is", "Implied bookings and backlog for every US filer that discloses RPO (remaining performance obligation) - 1,696 filers, against the 22 in the semis study. Built from the SEC XBRL frames API: 309 requests for the whole market, because frames return every filer reporting a concept for a period in one call."),
            ("The core identity", "implied gross bookings = revenue + change in RPO.   book-to-bill = implied bookings / revenue.   Above 1.0 means orders are running ahead of what is being recognised; below 1.0 means revenue is being met out of backlog."),
            ("The gate that matters", "RPO cover (RPO / quarterly revenue) must exceed 1 quarter. Where cover is near zero - NVIDIA, AMD, Lam - the derived book-to-bill pins to 1.00 and carries no information, so those names would inject pure noise into a cross-sectional rank. This gate was found empirically, not assumed."),
            ("backlog_ahead", "book-to-bill above 1.05 and rising while revenue growth is still under 15%. The pre-inflection state: contracted work piling up before it reaches the P&L. This is the screen worth watching."),
            ("accelerating", "RPO growth accelerating with book-to-bill above 1. Already inflecting, and confirmed by orders."),
            ("borrowing", "revenue growing while book-to-bill is below 1. Revenue is being met out of backlog rather than new orders - invisible if you look at revenue alone. ON Semi's pattern."),
            ("rpo_step", "single-quarter jumps in RPO over 50%. Usually a disclosure event - a scope change or first-time disclosure - not a quarter of orders, so these are held out of the acceleration rank rather than scored with it."),
            ("", ""),
            ("LIMITATION - restatement bias", "The frames API returns ONE observation per filer per period: the most recently filed. So this panel is latest-reported, not as-first-reported. That is correct for a current screen but means these signals CANNOT be backtested on this data - a return study has to go back through the per-filer path that records filing dates."),
            ("LIMITATION - coverage", "Frames occasionally skip a filer-quarter (median coverage 100%, 25th percentile 85%). A gap breaks the consecutive-quarter chain, so the change in RPO is left blank for that quarter rather than computed across a gap."),
            ("Data hygiene", "Three explicit checks, because a market-wide pull has no reconciliation validator behind it: RPO cover over 40 quarters is treated as a bad source fact and excluded (TPI Composites tags $59.6bn against $72m of quarterly revenue); the revenue concept is pinned per filer so it cannot switch mid-series; and four quarters are reconciled to the annual figure where the fiscal year ends in December, which is the only case where that test is valid."),
            ("Revenue concepts", "ASC 606 revenue is a SUBSET of total revenue, not an alternative to it - for a lessor like Crown Castle the 606 tag is 3% of the top line. One concept is pinned per filer: the one with the most quarters, and among equals the larger."),
            ("Validated", "Book-to-bill from this path reproduces the per-filer numbers exactly for AVGO, SNPS, CDNS and ON, and RPO matches companyfacts to the dollar for AKAM, LMT, DELL, PWR and XYL.")]
    sheet(wb, "How to read this", pd.DataFrame(rows, columns=["", " "]), widths={"": 26, " ": 130}, freeze_col=1)

    NOTE = {"backlog_ahead": "Book-to-bill above 1.05 and rising while revenue growth is still soft - contracted work building ahead of the P&L.",
            "accelerating": "RPO growth accelerating with book-to-bill above 1. Disclosure-step events are excluded and sit on their own tab.",
            "borrowing": "Revenue growing while book-to-bill is below 1: the top line is being met out of backlog, not new orders.",
            "rpo_step": "Single-quarter RPO jumps over 50%. Read these as disclosure events until the footnote says otherwise."}
    TITLE = {"backlog_ahead": "Backlog building ahead", "accelerating": "Accelerating backlog",
             "borrowing": "Revenue borrowing from backlog", "rpo_step": "RPO step - disclosure events"}
    for k, d in S.items():
        sheet(wb, TITLE[k], d[TS.COLS], note=NOTE[k], widths={"name": 34, "industry": 40},
              heat=["b2b", "b2b_trend", "rpo_yoy", "rpo_yoy_accel"])

    top = list(dict.fromkeys(list(S["backlog_ahead"].head(16).cik) + list(S["accelerating"].head(14).cik)))
    sheet(wb, "Trajectory", traj(P, top), note="RPO level ($m) and book-to-bill by quarter for the names at the top of the first two screens. The level tells you the size of the commitment; the book-to-bill series tells you whether it is still being added to.", widths={"name": 32})

    ind = G.groupby("industry").agg(names=("cik", "size"), median_b2b=("b2b", "median"), median_rpo_yoy=("rpo_yoy", "median"),
                                    median_cover=("rpo_cover_q", "median"), rising=("b2b_trend", lambda s: (s > 0).mean())).reset_index()
    ind = ind[ind.names >= 3].sort_values("median_b2b", ascending=False)
    sheet(wb, "By industry", ind, note="Where in the market backlog is building. Only industries with at least three gated filers. 'rising' is the share whose book-to-bill is above its own prior three quarters.",
          widths={"industry": 46}, heat=["median_b2b", "median_rpo_yoy", "rising"], freeze_col=2)

    semis = [int(x) for x in pd.read_pickle("out/rpo_facts.pkl").cik.unique()]
    sm = L[L.cik.isin(semis)].copy()
    sheet(wb, "Semis in market context", sm[TS.COLS + ["gated", "rpo_quarters"]].sort_values("b2b", ascending=False),
          note="The original semis universe scored on the same basis as the rest of the market, including the names the RPO-cover gate excludes - which is most of them.",
          widths={"name": 32, "industry": 38}, heat=["b2b", "rpo_yoy"])

    sheet(wb, "Full gated universe", G[TS.COLS].sort_values("b2b", ascending=False), note=f"All {len(G)} filers clearing the gate, ranked by book-to-bill.",
          widths={"name": 34, "industry": 40}, heat=["b2b", "rpo_yoy"])
    ex = L[~L.gated].copy()
    ex["excluded_because"] = np.where(~ex.clean, "failed a data-quality check",
                              np.where(ex.rpo_cover_q < TS.MIN_COVER, "RPO cover under 1 quarter - book-to-bill not meaningful",
                              np.where(ex.rev < TS.MIN_REV_Q, "revenue under $25m/quarter",
                              np.where(ex.cqp < TS.CURRENT, "no recent filing", "under 5 quarters of RPO history"))))
    sheet(wb, "Excluded and why", ex[["ticker", "name", "cqp", "rev", "rpo", "rpo_cover_q", "b2b", "rpo_quarters", "excluded_because"]].sort_values("rev", ascending=False),
          note="Every filer the gate dropped, with the reason - so nothing is silently missing.", widths={"name": 34, "excluded_because": 52})

    t2 = G[(G.b2b > 1.05) | (G.rpo_yoy_accel > 0.10)].sort_values("b2b", ascending=False)
    sheet(wb, "Tier 2 shortlist", t2[["ticker", "name", "industry", "b2b", "b2b_trend", "rpo_yoy", "rpo_yoy_accel", "rev", "rpo_cover_q"]],
          note=f"The {len(t2)} names worth running the deep segment-level pass on. Tier 2 costs about 335 SEC requests and 47MB per name, so this list is what keeps it to a few hours rather than a few days.",
          widths={"name": 34, "industry": 40}, heat=["b2b", "rpo_yoy"])

    src = pd.DataFrame([("SEC XBRL frames API", f"{len(pd.read_pickle('out/tier1_fetchlog.pkl')) if os.path.exists('out/tier1_fetchlog.pkl') else 309} requests covering 2017Q1-{pd.Timestamp.now(tz='UTC').tz_localize(None).to_period('Q')}: RPO and contract liabilities as instant frames, revenue as duration frames plus annual frames for deriving Q4"),
                        ("Concepts", "RevenueRemainingPerformanceObligation; ContractWithCustomerLiability and ...Current; RevenueFromContractWithCustomer{Ex,In}cludingAssessedTax and Revenues"),
                        ("SEC company_tickers.json", "CIK to ticker; blanks are filers with no listed ticker (subsidiaries, funds)"),
                        ("SEC submissions API", "SIC industry for the gated shortlist only - one request each, so it does not run over all 8,447 filers"),
                        ("Scale", "all figures $m. 173,233 filer-quarters, 8,447 filers, 1,696 of them disclosing RPO"),
                        ("What tier 1 cannot do", "no segment detail - frames exclude dimensional facts, so segment and product lines need the tier-2 per-filer parse")],
                       columns=["source", "detail"])
    ws = sheet(wb, "Method and sources", src, widths={"source": 26, "detail": 120}, freeze_col=1)
    for c in ws["B"]: c.alignment = WRAP
    wb.save("out/market_bookings_screen.xlsx")
    print("wrote out/market_bookings_screen.xlsx |", len(wb.sheetnames), "sheets")
    print(wb.sheetnames)
if __name__ == "__main__": main()
