"""SMH + SOXX universe workbook: current issuer holdings (24 Sep 2026), SEC N-PORT holdings (30 Jun 2026), mapped universe with SEC/Fiscal coverage."""
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from build_workbooks import title, hdr, put, table, F_B, F_BB, F_I, F_BL
SMH_CUR = """NVDA|Nvidia Corp|19.28
TSM|Taiwan Semiconductor Manufacturing Co L|9.26
AMD|Advanced Micro Devices Inc|5.80
AVGO|Broadcom Inc|5.25
INTC|Intel Corp|5.17
MU|Micron Technology Inc|5.05
SKHY|Sk Hynix Inc|4.50
TXN|Texas Instruments Inc|4.46
MRVL|Marvell Technology Inc|4.44
KLAC|Kla Corp|4.41
AMAT|Applied Materials Inc|4.36
ASML|Asml Holding Nv|4.30
LRCX|Lam Research Corp|4.19
QCOM|Qualcomm Inc|3.98
ADI|Analog Devices Inc|3.96
CDNS|Cadence Design Systems Inc|1.89
SNPS|Synopsys Inc|1.76
TER|Teradyne Inc|1.23
ARM|Arm Holdings Plc|1.21
MPWR|Monolithic Power Systems Inc|1.19
ALAB|Astera Labs Inc|1.15
NXPI|Nxp Semiconductors Nv|0.99
STM|Stmicroelectronics Nv|0.89
MCHP|Microchip Technology Inc|0.76
ON|On Semiconductor Corp|0.46
USD CASH|Cash|0.02"""
SOXX_CUR = """INTC|INTEL CORPORATION|10.20
AMD|ADVANCED MICRO DEVICES|9.58
MU|MICRON TECHNOLOGY|8.08
NVDA|NVIDIA|7.29
AVGO|BROADCOM INC|6.78
MRVL|MARVELL TECHNOLOGY|4.39
QCOM|QUALCOMM|4.08
KLAC|KLA|3.82
ADI|ANALOG DEVICES|3.79
TXN|TEXAS INSTRUMENT INC|3.72
AMAT|APPLIED MATERIAL INC|3.71
LRCX|LAM RESEARCH|3.65
MPWR|MONOLITHIC POWER SYSTEMS INC|3.14
SKHY|SK HYNIX AMERICAN DEPOSITARY SHARE|3.11
TER|TERADYNE|2.98
TSM|TAIWAN SEMICONDUCTOR MANUFACTURING|2.98
NXPI|NXP SEMICONDUCTORS NV|2.87
ALAB|ASTERA LABS|2.54
MCHP|MICROCHIP TECHNOLOGY INC|1.97
CRDO|CREDO TECHNOLOGY GROUP HOLDING LTD|1.64
ON|ON SEMICONDUCTOR CORP|1.40
ASML|ASML HOLDING ADR REPRESENTING|1.39
TSEM|TOWER SEMICONDUCTOR|1.13
ENTG|ENTEGRIS INC|1.12
MTSI|MACOM TECHNOLOGY SOLUTIONS|1.00
CBRS|CEREBRAS SYSTEMS INC CLASS A|0.98
ASX|ASE TECHNOLOGY HOLDING ADR REPRESE|0.88
UMC|UNITED MICRO ELECTRONICS ADR REP|0.67
ARM|ARM HOLDINGS AMERICAN DEPOSITARY S|0.53
STM|STMICROELECTRONICS ADR|0.40
XTSLA|BLK CSH FND TREASURY SL AGENCY|0.09
USD|USD CASH|0.06
WFFUT|CASH COLLATERAL USD WFFUT|0.01"""
CASH = {"USD CASH", "USD", "XTSLA", "WFFUT"}
def parse(s, etf, src): return pd.DataFrame([dict(etf=etf, ticker=a, name=b, weight_pct=float(c), as_of="2026-09-24", source=src) for a, b, c in (l.split("|") for l in s.splitlines())])
cur = pd.concat([parse(SMH_CUR, "SMH", "vaneck.com holdings page"), parse(SOXX_CUR, "SOXX", "ishares.com latest-holdings.csv")])
npt = pd.read_pickle("out/etf_holdings_nport.pkl"); u = pd.read_pickle("out/etf_universe_map.pkl")
npt = npt.merge(u[["isin", "ticker"]], on="isin", how="left")
FISCAL = {t: ("NYSE_" if t in ("TSM", "STM", "ASX", "UMC") else "NASDAQ_") + t for t in set(cur.ticker) | set(u.ticker.dropna())}
DATASETS = {"NVMI": "financials, adjusted, prices (no segments/KPIs)"}
CCY = {"TSM": "TWD", "ASX": "TWD", "UMC": "TWD", "ASML": "EUR", "SKHY": "KRW"}
extra = pd.DataFrame([dict(ticker="SKHY", name="SK hynix Inc", cik=2120882, annual_form="none yet (6-K only; ADR since 2026)", files_10q=False, first_annual=None, country="KR"),
                      dict(ticker="TSEM", name="Tower Semiconductor Ltd", cik=928876, annual_form="20-F", files_10q=False, first_annual="1996 (20-F series)", country="IL"),
                      dict(ticker="CBRS", name="Cerebras Systems Inc", cik=2021728, annual_form="none yet (10-Q x2)", files_10q=True, first_annual=None, country="US")])
uni = pd.concat([u[["ticker", "name", "cik", "annual_form", "files_10q", "first_annual", "country"]], extra], ignore_index=True)
eqc = cur[~cur.ticker.isin(CASH)]
for e in ("SMH", "SOXX"):
    uni[f"{e} wt % (24 Sep)"] = uni.ticker.map(dict(zip(eqc[eqc.etf == e].ticker, eqc[eqc.etf == e].weight_pct)))
    n = npt[(npt.etf == e) & (npt.asset_cat == "EC")]; uni[f"{e} wt % (30 Jun)"] = uni.ticker.map(dict(zip(n.ticker, n.weight_pct)))
DONE = {"NVDA": "done (NVDA workbook + warehouse)", "AMD": "done (step 2)", "TXN": "done (step 2)", "ASML": "done (step 2)"}
def path(r):
    t = r.ticker
    if t == "SKHY": return "no SEC history: needs Korean DART filings (or Fiscal.ai KRW data)"
    if t in ("AVGO", "MRVL"): return "10-K path; pre-redomicile history under predecessor CIK (Avago/Broadcom Ltd; Marvell Technology Group) - add a CIK chain"
    if "20-F" in str(r.annual_form): return "FPI path (20-F annual + 6-K results releases), like ASML" + ("; TWD reporting" if CCY.get(t) == "TWD" else "")
    if t in ("CBRS", "ALAB", "CRDO", "ARM"): return "10-K path; short public history (IPO " + {"CBRS": "2025-26", "ALAB": "2024", "CRDO": "2022", "ARM": "2023"}[t] + ")"
    return "10-K path (same as AMD/TXN)"
uni["pipeline status"] = uni.ticker.map(DONE).fillna("to do"); uni["pipeline path"] = uni.apply(path, axis=1)
uni["Fiscal.ai key"] = uni.ticker.map(FISCAL); uni["Fiscal.ai datasets"] = uni.ticker.map(DATASETS).fillna("financials, segments & KPIs, adjusted, prices")
uni["reporting ccy"] = uni.ticker.map(CCY).fillna("USD")
uni = uni.sort_values(["SMH wt % (24 Sep)", "SOXX wt % (24 Sep)"], ascending=False, na_position="last")

wb = Workbook(); ws = wb.active; ws.title = "Summary"
title(ws, "SMH + SOXX holdings and the pipeline universe", "Fiscal.ai does not carry ETFs or ETF holdings (checked: no SMH/SOXX instruments; fund data is manager-level 13F only). Holdings come from the issuers (current) and SEC Form N-PORT (quarter-end, filed ~60 days later).")
cur_in = set(eqc.ticker); jun_in = set(npt[npt.asset_cat == "EC"].ticker.dropna())
rows = [("Current holdings", "24 Sep 2026 — VanEck page (SMH, 25 stocks) and iShares CSV (SOXX, 30 stocks)"),
        ("Quarter-end holdings", "30 Jun 2026 — SEC N-PORT: SMH 25 stocks (filed 26 Aug), SOXX 30 stocks + futures + cash (filed 25 Aug); shares, value, CUSIP, ISIN"),
        ("Universe (union of all)", f"{len(uni)} companies: {len(cur_in)} held today, {len(jun_in - cur_in)} held in June but since dropped"),
        ("Added since June", ", ".join(sorted(cur_in - jun_in)) + "  (SK Hynix in both ETFs; Tower and Cerebras in SOXX)"),
        ("Dropped since June", ", ".join(sorted(jun_in - cur_in)) + "  (Skyworks from both; Nova, Rambus from SOXX)"),
        ("In both ETFs today", f"{len(set(eqc[eqc.etf == 'SMH'].ticker) & set(eqc[eqc.etf == 'SOXX'].ticker))} companies"),
        ("Pipeline status", "done: NVDA, AMD, TXN, ASML; to do: " + str(int((uni['pipeline status'] == 'to do').sum()))),
        ("Pipeline paths", "; ".join(f"{k}: {v}" for k, v in uni['pipeline path'].str.split(r"[;(]").str[0].str.strip().value_counts().items())),
        ("Watch-outs", "SK Hynix has no SEC history (Korean DART / Fiscal.ai KRW needed); AVGO and MRVL need predecessor CIKs; TSM/ASX/UMC report in TWD; CBRS/ALAB/CRDO/ARM have short histories.")]
r = 4; hdr(ws, r, 1, "Item", 26); hdr(ws, r, 2, "Detail", 120); r += 1
for a, b in rows: put(ws, r, 1, a, font=F_BB); x = put(ws, r, 2, b); x.alignment = Alignment(wrap_text=True, vertical="top"); r += 1
ws2 = wb.create_sheet("Universe"); title(ws2, "Universe: every company held by SMH or SOXX (today or at 30 Jun 2026), with SEC and Fiscal.ai coverage",
    "Weights are hardcoded from the sources (blue). 'Max weight' and 'held by' are formulas. Sorted by current SMH weight.")
cols = ["ticker", "name", "SMH wt % (24 Sep)", "SOXX wt % (24 Sep)", "SMH wt % (30 Jun)", "SOXX wt % (30 Jun)", "cik", "annual_form", "files_10q", "first_annual", "country",
        "reporting ccy", "Fiscal.ai key", "Fiscal.ai datasets", "pipeline status", "pipeline path"]
end = table(ws2, uni[cols].rename(columns={"cik": "SEC CIK", "annual_form": "annual form", "files_10q": "files 10-Q", "first_annual": "first annual filing"}), 4,
            widths={"name": 34, "pipeline path": 70, "Fiscal.ai datasets": 38, "annual form": 24, "pipeline status": 26, "Fiscal.ai key": 16}, nfs={c: "0.00" for c in cols if "wt" in c} | {"SEC CIK": "0"})
c1 = len(cols) + 1; hdr(ws2, 4, c1, "Held by (today)", 16); hdr(ws2, 4, c1 + 1, "Max weight today %", 14)
for i in range(5, end):
    ws2.cell(i, c1, f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(D{i})),"both",IF(ISNUMBER(C{i}),"SMH only",IF(ISNUMBER(D{i}),"SOXX only","dropped since June")))').font = F_B
    x = ws2.cell(i, c1 + 1, f'=IF(COUNT(C{i}:D{i})=0,"",MAX(C{i}:D{i}))'); x.number_format = "0.00"; x.font = F_B
    for j in (3, 4, 5, 6): ws2.cell(i, j).font = F_BL
put(ws2, end + 1, 1, "Totals (formulas):", font=F_BB)
for j in (3, 4, 5, 6): x = ws2.cell(end + 1, j, f"=SUM({L(j)}5:{L(j)}{end - 1})"); x.number_format = "0.00"; x.font = F_BB
ws3 = wb.create_sheet("Current holdings"); title(ws3, "Current holdings as published by the issuers, 24 Sep 2026 (cash lines included)", "SMH: vaneck.com/us/en/investments/semiconductor-etf-smh/holdings/ ; SOXX: ishares.com .../latest-holdings.csv")
table(ws3, cur, 4, widths={"name": 40, "source": 30}, nfs={"weight_pct": "0.00"})
ws4 = wb.create_sheet("N-PORT 30 Jun 2026"); title(ws4, "Full portfolio from SEC Form N-PORT (report date 30 Jun 2026)", "SOXX: iShares Trust, series S000004354, accession 0002071691-26-019781. SMH: VanEck ETF Trust, series S000034411, accession 0001410368-26-086937.")
n = npt[["etf", "ticker", "name", "title", "cusip", "isin", "shares", "value_usd", "weight_pct", "asset_cat", "country", "report_date", "filed", "accn"]].sort_values(["etf", "weight_pct"], ascending=[True, False])
table(ws4, n, 4, widths={"name": 38, "title": 40, "accn": 22}, nfs={"shares": "#,##0", "value_usd": "#,##0", "weight_pct": "0.000"})
ws5 = wb.create_sheet("Changes since June"); title(ws5, "Membership changes between N-PORT (30 Jun) and issuer holdings (24 Sep)")
ch = []
for e in ("SMH", "SOXX"):
    a = set(eqc[eqc.etf == e].ticker); b = set(npt[(npt.etf == e) & (npt.asset_cat == "EC")].ticker.dropna())
    ch += [dict(etf=e, change="added", ticker=t, weight_now=float(eqc[(eqc.etf == e) & (eqc.ticker == t)].weight_pct.iloc[0]), weight_june=None) for t in sorted(a - b)]
    ch += [dict(etf=e, change="dropped", ticker=t, weight_now=None, weight_june=float(npt[(npt.etf == e) & (npt.ticker == t)].weight_pct.iloc[0])) for t in sorted(b - a)]
table(ws5, pd.DataFrame(ch), 4, nfs={"weight_now": "0.00", "weight_june": "0.00"})
wb.save("out/semis_etf_universe.xlsx"); print("saved", len(uni), "companies")
print(uni[["ticker", "SMH wt % (24 Sep)", "SOXX wt % (24 Sep)", "pipeline status"]].to_string(index=False))
