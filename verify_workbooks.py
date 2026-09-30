"""Read the recalculated workbooks back and reconcile against the validated pipeline outputs + fixed spot values from filings."""
import openpyxl, pandas as pd, numpy as np, json, tables as T, sys
def grid_read(ws, r0=4, fc=3):
    dates = {c: ws.cell(r0 + 1, c).value for c in range(fc, ws.max_column + 1) if ws.cell(r0 + 1, c).value}
    out = {}
    for r in range(r0 + 2, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None: continue
        if lab in out: continue                     # labels repeat per section (revenue first, then operating income): keep the first
        out[lab] = {pd.Timestamp(d): ws.cell(r, c).value for c, d in dates.items() if ws.cell(r, c).value not in (None, "")}
    return out
res = []
def ok(t, what, cond, detail=""): res.append((t, what, "PASS" if cond else "FAIL", detail))
for t in sys.argv[1:]:
    wb = openpyxl.load_workbook(f"out/{t}_sec_history.xlsx", data_only=True); cfg = json.load(open(f"config/{t}.json")); fpi = cfg.get("filer_type") == "FPI"
    g = grid_read(wb["Income stmt quarterly"])
    if fpi:
        q = pd.read_pickle("out/ASML_q.pkl"); src = dict(zip(pd.to_datetime(q.period_end), q.revenue))
    else:
        f = pd.read_pickle(f"out/{t}_fin.pkl"); s = f[(f.canonical == "revenue") & (f.duration_months == 3) & (f.basis == "as_first_reported")]; src = dict(zip(pd.to_datetime(s.period_end), s.value))
    rv = g["Revenue"]; mism = [p for p in src if abs(rv.get(p, np.nan) - src[p]) > 1e-6]
    ok(t, "IS quarterly revenue = source (all quarters)", not mism and len(rv) == len(src), f"{len(rv)} quarters, mismatches {mism[:3]}")
    gm = g["Gross margin %"]; gp = g["Gross profit"]; bad = [p for p in gm if abs(gm[p] - gp[p] / rv[p]) > 1e-9]
    ok(t, "Gross margin formula = GP / revenue", not bad, f"{len(gm)} cells")
    yoy = g["Revenue growth y/y %"]; ps = sorted(rv); bad = [p for i, p in enumerate(ps) if i >= 4 and p in yoy and abs(yoy[p] - (rv[p] / rv[ps[i - 4]] - 1)) > 1e-9]
    ok(t, "Revenue y/y formula points 4 columns back", not bad and len(yoy) >= len(ps) - 5, f"{len(yoy)} cells")
    if not fpi:
        seg, sub, geo = T.segments(t); dq = T.dim_quarterly(seg, ["name"])
        gs = grid_read(wb["Segments quarterly"]); n = bad = 0
        for (nm, met), s in dq.groupby(["name", "metric"]):
            # segment rows appear under the metric section in the same order; map by label within section
            pass
        ws = wb["Segments quarterly"]; sec = None; rows = {}
        for r in range(6, ws.max_row + 1):
            lab = ws.cell(r, 1).value
            if lab in ("Segment revenue", "Segment operating income"): sec = "revenue" if lab == "Segment revenue" else "operating_income"; continue
            if lab and sec and ws.cell(r, 2).value and "share" not in lab and "y/y" not in lab and "within" not in lab: rows[(sec, lab)] = r
            if lab and lab.startswith("Revenue within"): sec = None
        dates = {c: pd.Timestamp(ws.cell(5, c).value) for c in range(3, ws.max_column + 1) if ws.cell(5, c).value}
        for (met, nm), r in rows.items():
            s = dq[(dq.metric == met) & (dq.name == nm)]
            for p, v in zip(s.period_end, s["first"]):
                c = [k for k, d in dates.items() if d == p]
                if not c: continue
                n += 1; x = ws.cell(r, c[0]).value
                if x is None or abs(x - v) > 1e-6: bad += 1
        ok(t, "Segment cells = validated segment table", n > 100 and bad == 0, f"{n} cells checked, {bad} off")
    gv = wb["Guidance vs actual"]; ck = pd.read_pickle(f"out/{t}_guid_check.pkl"); ck["accn"] = ck.accn.astype(str)
    xs = {gv.cell(r, 2).value: gv.cell(r, 12).value for r in range(5, gv.max_row + 1) if gv.cell(r, 2).value}
    m = ck[ck.status == "ok"]; bad = [a for a, rr in zip(m.accn, m.ratio) if a in xs and isinstance(xs[a], float) and abs(xs[a] - (rr - 1)) > 1e-6]
    ok(t, "Guidance beat % formula = validator ratio", not bad and sum(a in xs for a in m.accn) >= 0.9 * len(m), f"{sum(a in xs for a in m.accn)} of {len(m)} ok-rows present, {len(bad)} differ")
    wb2 = openpyxl.load_workbook(f"out/{t}_sec_history.xlsx", data_only=True)
    errs = sum(1 for ws in wb2.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#"))
    ok(t, "No error strings anywhere", errs == 0, str(errs))
# fixed spot checks against filings (values read from the filing text earlier in this session)
W = {k: openpyxl.load_workbook(f"out/{k}_sec_history.xlsx", data_only=True) for k in sys.argv[1:]}
def cell(t, sheet, label, date):
    g = grid_read(W[t][sheet]); return g.get(label, {}).get(pd.Timestamp(date))
spots = []
if "AMD" in W: spots += [("AMD", "Segments quarterly", "Data Center", "2023-09-30", 1598, "10-Q table row 'Data Center 1,598' (XBRL tag said Client)"),
                         ("AMD", "Segments quarterly", "Xilinx", "2022-03-26", 559, "Q1 2022 10-Q segment table"),
                         ("AMD", "Segments quarterly", "Gaming (within Client and Gaming)", "2025-12-27", 843, "FY2025 10-K minus 9M; matches Fiscal.ai")]
if "TXN" in W: spots += [("TXN", "End markets annual", "Industrial", "2025-12-31", 33, "10-K 2025 chart (33% of TI revenue)"),
                         ("TXN", "Segments quarterly", "Analog", "2012-06-30", 1800, "label linkbase rule (ReportableSegmentsMember = Analog)")]
if "ASML" in W: spots += [("ASML", "KPIs quarterly", "Net bookings", "2025-12-31", 13158, "Q4 2025 release headline table"),
                          ("ASML", "KPIs quarterly", "Net bookings", "2021-06-30", 8271, "Q2 2021 headline table (text said 'EUR 8.3 billion')"),
                          ("ASML", "Income stmt quarterly", "Net income", "2013-06-30", 220.8, "first printing (later restated to 245.1)"),
                          ("ASML", "Income stmt latest", "Net income", "2013-06-30", 245.1, "restated in the Q4 2013 release"),
                          ("ASML", "Income stmt quarterly", "Net income", "2026-03-31", 2756.7, "PDF-flattened full-precision table")]
for t, sh, lab, d, exp, why in spots:
    v = cell(t, sh, lab, d); ok(t, f"spot: {sh} / {lab} / {d} = {exp}", v is not None and abs(v - exp) < 0.05, f"got {v}; {why}")
b = openpyxl.load_workbook("out/ASML_sec_history.xlsx", data_only=True) if "ASML" in W else None
if b:
    g = grid_read(b["Sales breakdown annual"]); ok("ASML", "spot: 2025 EUV (NXE) = 10,446", abs(g["EUV (NXE, low-NA)"][pd.Timestamp("2025-12-31")] - 10446) < 1, "20-F FY2025 XBRL")
    btb = grid_read(b["KPIs quarterly"])["Book-to-bill (net bookings / net system sales)"]; ok("ASML", "book-to-bill within 0.2-3.5", all(0.2 < v < 3.5 for v in btb.values()), f"{len(btb)} quarters, range {min(btb.values()):.2f}-{max(btb.values()):.2f}")
if "TXN" in W:
    s = grid_read(W["TXN"]["End markets annual"])["Sum of shares (check)"]; ok("TXN", "end-market shares sum 94-104", all(94 <= v <= 104 for v in s.values()), str(sorted(set(s.values()))))
r = pd.DataFrame(res, columns=["ticker", "check", "result", "detail"]); pd.set_option("display.width", 250); pd.set_option("display.max_colwidth", 90)
print(r.to_string(index=False)); print(r.result.value_counts().to_dict())
