"""Read every workbook back and reconcile it against the validated pipeline outputs."""
import openpyxl, pandas as pd, numpy as np, json, os, glob, sys, tables as T
from build_universe_data import qrev
def grid_read(ws, r0=4, fc=3):
    dates = {c: ws.cell(r0 + 1, c).value for c in range(fc, ws.max_column + 1) if ws.cell(r0 + 1, c).value}
    out = {}
    for r in range(r0 + 2, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None or lab in out: continue
        out[lab] = {pd.Timestamp(d): ws.cell(r, c).value for c, d in dates.items() if ws.cell(r, c).value not in (None, "")}
    return out
res = []
for t in sorted({os.path.basename(p).split("_sec_history")[0] for p in glob.glob("out/*_sec_history.xlsx")}):
    cfg = json.load(open(f"config/{t}.json")); wb = openpyxl.load_workbook(f"out/{t}_sec_history.xlsx", data_only=True)
    r = dict(ticker=t, tabs=len(wb.worksheets))
    errs = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#"))
    r["error_cells"] = errs
    if "Income stmt quarterly" in wb.sheetnames:
        g = grid_read(wb["Income stmt quarterly"]); src = qrev(t, cfg); rv = g.get("Revenue", {})
        miss = [p for p in src.index if abs(rv.get(p, np.nan) - src[p]) > 1e-6]
        r.update(quarters_in_book=len(rv), quarters_in_source=len(src), revenue_mismatches=len(miss))
        gm, gp = g.get("Gross margin %", {}), g.get("Gross profit", {})
        r["gm_formula_bad"] = sum(1 for p in gm if p in gp and p in rv and rv[p] and abs(gm[p] - gp[p] / rv[p]) > 1e-9)
    if "Segments quarterly" in wb.sheetnames:
        seg, sub, geo = T.segments(t); dq = T.dim_quarterly(seg, ["name"])
        ws = wb["Segments quarterly"]; dates = {c: pd.Timestamp(ws.cell(5, c).value) for c in range(3, ws.max_column + 1) if ws.cell(5, c).value}
        sec = None; rows = {}
        for rr in range(6, ws.max_row + 1):
            lab = ws.cell(rr, 1).value
            if lab in ("Segment revenue", "Segment operating income"): sec = "revenue" if lab == "Segment revenue" else "operating_income"; continue
            if lab and str(lab).startswith(("Revenue within", "Segment share", "Segment revenue growth")): sec = None; continue
            if lab and sec and ws.cell(rr, 2).value: rows[(sec, lab)] = rr
        n = bad = 0
        for (met, nm), rr in rows.items():
            s = dq[(dq.metric == met) & (dq.name == nm)]
            for p, v in zip(s.period_end, s["first"]):
                c = [k for k, d in dates.items() if d == p]
                if not c: continue
                n += 1; x = ws.cell(rr, c[0]).value
                if x is None or abs(x - v) > 1e-6: bad += 1
        r.update(segment_cells=n, segment_bad=bad)
    if "Guidance vs actual" in wb.sheetnames and os.path.exists(f"out/{t}_guid_check.pkl"):
        gv = wb["Guidance vs actual"]; ck = pd.read_pickle(f"out/{t}_guid_check.pkl")
        xs = {gv.cell(rr, 2).value: gv.cell(rr, 12).value for rr in range(5, gv.max_row + 1) if gv.cell(rr, 2).value}
        m = ck[ck.status == "ok"]; bad = [a for a, rt in zip(m.accn, m.ratio) if a in xs and isinstance(xs[a], float) and abs(xs[a] - (rt - 1)) > 1e-6]
        r.update(guide_rows_checked=int(sum(a in xs for a in m.accn)), guide_formula_bad=len(bad))
    res.append(r)
d = pd.DataFrame(res); pd.set_option("display.width", 220)
print(d.fillna("").to_string(index=False))
tot = {c: int(d[c].sum()) for c in ("error_cells", "revenue_mismatches", "gm_formula_bad", "segment_bad", "guide_formula_bad") if c in d}
print("\nTOTAL PROBLEMS:", tot)
d.to_pickle("out/verify_all.pkl")
