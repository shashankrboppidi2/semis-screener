"""One Excel workbook per ticker from the validated pipeline outputs. Values are hardcoded (blue); ratios/growth/beat are live formulas (black).
Shading: derived Q4 (= FY - 9M) light blue; model-read values light orange; restated/latest views on their own tabs."""
import json, os, re, sys, pandas as pd, numpy as np, edgar, tables as T
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

A = "Arial"; F_H = Font(name=A, bold=True, color="FFFFFF"); FILL_H = PatternFill("solid", fgColor="1F3864"); F_B = Font(name=A, size=10)
F_BL = Font(name=A, size=10, color="0000FF"); F_BB = Font(name=A, size=10, bold=True); F_T = Font(name=A, bold=True, size=14); F_I = Font(name=A, italic=True, size=9, color="595959")
F_G = Font(name=A, size=10, color="008000")
FILL_Q4 = PatternFill("solid", fgColor="DDEBF7"); FILL_M = PatternFill("solid", fgColor="FCE4D6"); FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
NAMES = {"revenue": "Revenue", "cost_of_revenue": "Cost of revenue", "gross_profit": "Gross profit", "research_and_development": "R&D", "rd": "R&D", "sga": "SG&A",
         "total_operating_expenses": "Total operating expenses", "operating_income": "Operating income", "income_before_tax": "Income before tax", "income_tax": "Income tax",
         "net_income": "Net income", "eps_basic": "EPS basic", "eps_diluted": "EPS diluted", "shares_basic": "Shares basic (M)", "shares_diluted": "Shares diluted (M)",
         "cash_and_equivalents": "Cash & equivalents", "marketable_securities": "Marketable securities", "accounts_receivable": "Accounts receivable", "inventories": "Inventories",
         "total_current_assets": "Total current assets", "total_assets": "Total assets", "accounts_payable": "Accounts payable", "total_current_liabilities": "Total current liabilities",
         "total_liabilities": "Total liabilities", "total_equity": "Total equity", "system_sales": "Net system sales", "service_sales": "Net service & field option sales"}
IS_ORDER = ["revenue", "system_sales", "service_sales", "cost_of_revenue", "gross_profit", "research_and_development", "rd", "sga", "total_operating_expenses", "operating_income",
            "income_before_tax", "income_tax", "net_income", "eps_basic", "eps_diluted", "shares_basic", "shares_diluted"]
BS_ORDER = ["cash_and_equivalents", "marketable_securities", "accounts_receivable", "inventories", "total_current_assets", "total_assets", "accounts_payable",
            "total_current_liabilities", "total_liabilities", "total_equity"]

def title(ws, text, note=None):
    ws["A1"] = text; ws["A1"].font = F_T
    if note: ws["A2"] = note; ws["A2"].font = F_I
def hdr(ws, r, c, v, w=None):
    x = ws.cell(r, c, v); x.font = F_H; x.fill = FILL_H; x.alignment = Alignment(wrap_text=True, vertical="center")
    if w: ws.column_dimensions[L(c)].width = w
    return x
def put(ws, r, c, v, nf=None, font=F_B, fill=None):
    if isinstance(v, (np.floating, np.integer)): v = v.item()
    if isinstance(v, float) and np.isnan(v): v = None
    if isinstance(v, pd.Timestamp): v = v.to_pydatetime(); nf = nf or "yyyy-mm-dd"
    x = ws.cell(r, c, v); x.font = font
    if nf: x.number_format = nf
    if fill: x.fill = fill
    return x
def table(ws, df, r0, widths=None, nfs=None, fills=None):
    """long table; nfs: {col: number_format}; fills: Series of fills per row (or None)"""
    for j, c in enumerate(df.columns, 1): hdr(ws, r0, j, c, (widths or {}).get(c, max(11, min(40, len(str(c)) + 2))))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        f = fills[i - r0 - 1] if fills is not None else None
        for j, v in enumerate(row, 1):
            col = df.columns[j - 1]; nf = (nfs or {}).get(col)
            if nf is None and isinstance(v, (float, np.floating)): nf = "#,##0.0;(#,##0.0);-"
            put(ws, i, j, v, nf, fill=f)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return r0 + len(df) + 1

def qlabel(pe, fye): fy, q = T.fq(pe, fye); return f"Q{q} FY{fy}" if fye != 12 else f"Q{q} {fy}"

def grid(ws, r0, periods, rows, fye, first_col=3, money="#,##0;(#,##0);-", label_w=38):
    """rows: list of dicts {label, values: {period: v}, flags: {period: 'q4'|'model'}, nf, kind: 'data'|'formula', formula: callable(col_letter, rowmap, j)}
    returns {key: row_number}"""
    hdr(ws, r0, 1, "Line item", label_w); hdr(ws, r0, 2, "Units / note", 16)
    hdr(ws, r0 + 1, 1, ""); hdr(ws, r0 + 1, 2, "period end →")
    for j, p in enumerate(periods, first_col):
        hdr(ws, r0, j, qlabel(p, fye) if fye else str(pd.Timestamp(p).year), 11); x = hdr(ws, r0 + 1, j, pd.Timestamp(p).to_pydatetime()); x.number_format = "yyyy-mm-dd"
    rowmap = {}; r = r0 + 2
    for spec in rows:
        if spec.get("kind") == "section":
            put(ws, r, 1, spec["label"], font=F_BB); r += 1; continue
        rowmap[spec["key"]] = r; put(ws, r, 1, spec["label"], font=F_BB if spec.get("bold") else F_B); put(ws, r, 2, spec.get("note", ""), font=F_I)
        for j, p in enumerate(periods, first_col):
            col = L(j)
            if spec.get("kind") == "formula":
                f = spec["formula"](col, rowmap, j, first_col)
                if f: put(ws, r, j, f, spec.get("nf", "0.0%"))
            else:
                v = spec["values"].get(p)
                if v is None or (isinstance(v, float) and np.isnan(v)): continue
                fl = spec.get("flags", {}).get(p); fill = FILL_Q4 if fl == "q4" else FILL_M if fl == "model" else None
                put(ws, r, j, float(v), spec.get("nf", money), font=F_BL, fill=fill)
        r += 1
    ws.freeze_panes = ws.cell(r0 + 2, first_col)
    return rowmap, r

def ratio(num, den):
    return lambda col, rm, j, fc: f'=IF(AND(ISNUMBER({col}{rm[num]}),ISNUMBER({col}{rm[den]}),{col}{rm[den]}<>0),{col}{rm[num]}/{col}{rm[den]},"")' if num in rm and den in rm else None
def growth(key, lag):
    def f(col, rm, j, fc):
        if key not in rm or j - lag < fc: return None
        pc = L(j - lag); r = rm[key]
        return f'=IF(AND(ISNUMBER({col}{r}),ISNUMBER({pc}{r}),{pc}{r}<>0),{col}{r}/{pc}{r}-1,"")'
    return f

def legend(ws, r):
    put(ws, r, 1, "Legend", font=F_BB)
    put(ws, r + 1, 1, "Blue number = value extracted from a filing (hardcoded). Black = live formula.", font=F_I)
    x = put(ws, r + 2, 1, "Derived Q4 = full year minus nine months (companies do not file a Q4 10-Q)", font=F_I); x.fill = FILL_Q4
    x = put(ws, r + 3, 1, "Model-read value (Haiku/Sonnet), accepted only after deterministic validation", font=F_I); x.fill = FILL_M
    return r + 5

def fin_long(t, cfg):
    """US filers: companyfacts. FPI: results releases (quarterly) + 20-F (annual/balance sheet); annual only when no results-table
    reader exists for that filer (IFRS foreign issuers)."""
    if cfg.get("filer_type") == "FPI" and not os.path.exists(f"out/{t}_q.pkl"):
        f = pd.read_pickle(f"out/{t}_fin.pkl"); f["period_end"] = pd.to_datetime(f.period_end)
        return f.iloc[0:0], f
    if cfg.get("filer_type") == "FPI":
        q = pd.read_pickle(f"out/{t}_q.pkl"); P = pd.read_pickle(f"out/{t}_q_printings.pkl")
        cols = [c for c in ["revenue", "system_sales", "service_sales", "gross_profit", "rd", "sga", "operating_income", "net_income"] if c in q]
        first = q.melt(id_vars=["period_end", "accn", "filed", "method"], value_vars=cols, var_name="canonical").dropna(subset=["value"]).assign(duration_months=3, basis="as_first_reported")
        P["filed"] = pd.to_datetime(P.filed); lat = P.sort_values("filed").melt(id_vars=["period_end", "accn", "filed"], value_vars=[c for c in cols if c in P], var_name="canonical").dropna(subset=["value"])
        lat = lat.sort_values("filed").drop_duplicates(["period_end", "canonical"], keep="last").assign(duration_months=3, basis="latest_filed", method="quarterly summary (later printing)")
        qq = pd.concat([first, lat], ignore_index=True); qq["period_end"] = pd.to_datetime(qq.period_end)
        for c in ("rd", "sga"): qq.loc[qq.canonical == c, "value"] = qq.loc[qq.canonical == c, "value"].abs()   # costs shown as positive
        f = pd.read_pickle(f"out/{t}_fin.pkl"); f["period_end"] = pd.to_datetime(f.period_end)
        return qq, f
    f = pd.read_pickle(f"out/{t}_fin.pkl"); f["period_end"] = pd.to_datetime(f.period_end); return f, f

def sheet_is(wb, t, cfg, cur, qd, basis, name, note):
    ws = wb.create_sheet(name); fye = cfg.get("fye_month", 12)
    d = qd[(qd.duration_months == 3) & (qd.basis == basis)]
    periods = sorted(d.period_end.unique()); money = "#,##0.0;(#,##0.0);-" if cfg.get("filer_type") == "FPI" else "#,##0;(#,##0);-"
    title(ws, f"{t} — income statement, quarterly ({'first reported' if basis == 'as_first_reported' else 'latest printing'}), {cur} millions", note)
    rows = []
    for c in IS_ORDER:
        s = d[d.canonical == c]
        if not len(s): continue
        vals = dict(zip(s.period_end, s.value))
        flags = {p: "q4" for p, cc in zip(s.period_end, s.get("concept", pd.Series([""] * len(s)))) if isinstance(cc, str) and "FY-9M" in cc}
        nf = "0.00" if c.startswith("eps") else ("#,##0.0" if c.startswith("shares") else money)
        rows.append(dict(key=c, label=NAMES.get(c, c), values=vals, flags=flags, nf=nf, bold=c in ("revenue", "gross_profit", "operating_income", "net_income"),
                         note={"eps_basic": cur + "/share", "eps_diluted": cur + "/share"}.get(c, cur + "m")))
    rows += [dict(kind="section", label="Ratios and growth (formulas)"),
             dict(key="gm", label="Gross margin %", kind="formula", formula=ratio("gross_profit", "revenue")),
             dict(key="om", label="Operating margin %", kind="formula", formula=ratio("operating_income", "revenue")),
             dict(key="nm", label="Net margin %", kind="formula", formula=ratio("net_income", "revenue")),
             dict(key="qoq", label="Revenue growth q/q %", kind="formula", formula=growth("revenue", 1)),
             dict(key="yoy", label="Revenue growth y/y %", kind="formula", formula=growth("revenue", 4))]
    rm, r = grid(ws, 4, periods, rows, fye, money=money); legend(ws, r + 1)
    return ws

def sheet_annual(wb, t, cfg, cur, f):
    ws = wb.create_sheet("Income stmt annual"); d = f[(f.duration_months == 12) & (f.basis == "as_first_reported")]
    periods = sorted(d.period_end.unique()); money = "#,##0.0;(#,##0.0);-" if cfg.get("filer_type") == "FPI" else "#,##0;(#,##0);-"
    title(ws, f"{t} — income statement, annual (as first reported in the {'20-F' if cfg.get('filer_type') == 'FPI' else '10-K'}), {cur} millions", "Source: SEC XBRL companyfacts; first filing that reported each year.")
    rows = [dict(key=c, label=NAMES.get(c, c), values=dict(zip(d[d.canonical == c].period_end, d[d.canonical == c].value)),
                 nf="0.00" if c.startswith("eps") else money, note=cur + "m") for c in IS_ORDER if len(d[d.canonical == c])]
    rows += [dict(kind="section", label="Ratios and growth (formulas)"), dict(key="gm", label="Gross margin %", kind="formula", formula=ratio("gross_profit", "revenue")),
             dict(key="om", label="Operating margin %", kind="formula", formula=ratio("operating_income", "revenue")), dict(key="yoy", label="Revenue growth y/y %", kind="formula", formula=growth("revenue", 1))]
    rm, r = grid(ws, 4, periods, rows, None, money=money)

def sheet_bs(wb, t, cfg, cur, f):
    ws = wb.create_sheet("Balance sheet"); d = f[(f.duration_months == 0) & (f.basis == "as_first_reported")]
    periods = sorted(d.period_end.unique()); money = "#,##0.0;(#,##0.0);-" if cfg.get("filer_type") == "FPI" else "#,##0;(#,##0);-"
    title(ws, f"{t} — balance sheet ({'annual, 20-F' if cfg.get('filer_type') == 'FPI' else 'quarter-end'}), as first reported, {cur} millions")
    rows = [dict(key=c, label=NAMES.get(c, c), values=dict(zip(d[d.canonical == c].period_end, d[d.canonical == c].value)), note=cur + "m") for c in BS_ORDER if len(d[d.canonical == c])]
    grid(ws, 4, periods, rows, cfg.get("fye_month", 12) if cfg.get("filer_type") != "FPI" else None, money=money)

def sheet_segments_annual(wb, t, cfg, cur, seg):
    """Filers that tag segments only on annual (and half-year) periods - IFRS foreign issuers - get an annual segment grid."""
    a = T.dim_annual(seg, ["name"])
    if not len(a): return {}
    ws = wb.create_sheet("Segments annual"); periods = sorted(a.period_end.unique())
    title(ws, f"{t} — reportable segments, annual (20-F XBRL), first reported, {cur} millions",
          "This filer tags segment figures on annual periods only, so there is no quarterly segment series. Names are the filer's own labels.")
    rows = []
    for met, lab in (("revenue", "Segment revenue"), ("operating_income", "Segment operating income")):
        s = a[a.metric == met]
        if not len(s): continue
        rows.append(dict(kind="section", label=lab))
        for nm, g in s.groupby("name"):
            rows.append(dict(key=f"{met}|{nm}", label=nm, values=dict(zip(g.period_end, g["first"])), note=cur + "m"))
    grid(ws, 4, periods, rows, None, money="#,##0.0;(#,##0.0);-")
    return {"segment_annual": a}
def sheet_segments(wb, t, cfg, cur):
    seg, sub, geo = T.segments(t); fye = cfg.get("fye_month", 12)
    out = {}
    for label, df, keys in (("segment", seg, ["name"]), ("sub-segment", sub, ["parent", "name"]), ("geography", geo, ["name"])):
        if len(df):
            q_ = T.dim_quarterly(df, keys)
            if len(q_): out[label] = q_
    if "segment" not in out: return sheet_segments_annual(wb, t, cfg, cur, seg)
    ws = wb.create_sheet("Segments quarterly"); q = out["segment"]
    model_keys = {(a, p) for a, p in zip(seg[seg.get("model", pd.Series(index=seg.index, dtype=object)).notna()].name, seg[seg.get("model", pd.Series(index=seg.index, dtype=object)).notna()].period_end)} if "model" in seg else set()
    periods = sorted(q.period_end.unique())
    title(ws, f"{t} — reportable segments, quarterly, first reported, {cur} millions",
          "Segment names change when the company re-segments; each scheme appears as its own rows. Latest (recast) values are on 'Segments detail'. "
          + ("AMD 10-Qs 2022-24 mis-tag Data Center/Client/Gaming in XBRL; values here follow the printed table rows (tag-swap guard)." if t == "AMD" else ""))
    rows = []
    fin = pd.read_pickle(f"out/{t}_fin.pkl"); fin["period_end"] = pd.to_datetime(fin.period_end)
    tot = fin[(fin.canonical == "revenue") & (fin.duration_months == 3) & (fin.basis == "as_first_reported")]
    rows.append(dict(key="total", label="Total revenue (consolidated)", values=dict(zip(tot.period_end, tot.value)), bold=True, note=cur + "m",
                     flags={p: "q4" for p, c in zip(tot.period_end, tot.concept) if "FY-9M" in c}))
    for met, mlabel in (("revenue", "Segment revenue"), ("operating_income", "Segment operating income")):
        rows.append(dict(kind="section", label=mlabel))
        names = q[q.metric == met].groupby("name").period_end.min().sort_values().index
        for nm in names:
            s = q[(q.metric == met) & (q.name == nm)]
            rows.append(dict(key=f"{met}|{nm}", label=nm, values=dict(zip(s.period_end, s["first"])), note=cur + "m",
                             flags={**{p: "q4" for p, src in zip(s.period_end, s.source) if src.startswith("derived")}, **{p: "model" for (n, p) in model_keys if n == nm}}))
    if "sub-segment" in out:
        rows.append(dict(kind="section", label="Revenue within segment (sub-segments)"))
        sq = out["sub-segment"]
        for (par, nm), s in sq[sq.metric == "revenue"].groupby(["parent", "name"]):
            rows.append(dict(key=f"sub|{nm}", label=f"{nm} (within {par})", values=dict(zip(s.period_end, s["first"])), note=cur + "m",
                             flags={p: "q4" for p, src in zip(s.period_end, s.source) if src.startswith("derived")}))
    rows.append(dict(kind="section", label="Segment share of total revenue (formulas)"))
    for nm in q[q.metric == "revenue"].groupby("name").period_end.min().sort_values().index:
        rows.append(dict(key=f"share|{nm}", label=f"{nm} share %", kind="formula", formula=ratio(f"revenue|{nm}", "total")))
    rows.append(dict(kind="section", label="Segment revenue growth y/y (formulas)"))
    for nm in q[q.metric == "revenue"].groupby("name").period_end.min().sort_values().index:
        rows.append(dict(key=f"yoy|{nm}", label=f"{nm} y/y %", kind="formula", formula=growth(f"revenue|{nm}", 4)))
    rm, r = grid(ws, 4, periods, rows, fye); legend(ws, r + 1)
    # detail (long) with first vs latest
    ws2 = wb.create_sheet("Segments detail")
    title(ws2, f"{t} — every segment value with first-reported and latest printing", "latest differs from first when the company recast a period (re-segmentation or restatement).")
    parts = []
    for label, qq in out.items():
        x = qq.copy(); x["dimension"] = label; x["segment"] = x.get("name"); x["parent"] = x["parent"] if "parent" in x else None
        parts.append(x)
    dd = pd.concat(parts)
    dd["quarter"] = [qlabel(p, fye) for p in dd.period_end]
    dd = dd[["dimension", "parent", "segment", "metric", "quarter", "period_end", "first", "latest", "source", "printings", "first_accn", "first_filed", "latest_accn", "latest_filed"]]
    dd = dd.sort_values(["dimension", "metric", "segment", "period_end"])
    table(ws2, dd, 4, widths={"segment": 34, "parent": 20, "first_accn": 22, "latest_accn": 22, "source": 22}, nfs={"first": "#,##0.0;(#,##0.0);-", "latest": "#,##0.0;(#,##0.0);-"})
    return out

def sheet_geo_annual(wb, t, cfg, cur, out):
    if "geography" not in out: return
    seg, sub, geo = T.segments(t); a = T.dim_annual(geo, ["name"])
    if not len(a): return
    ws = wb.create_sheet("Geography annual"); periods = sorted(a.period_end.unique())
    title(ws, f"{t} — revenue by geography, annual (XBRL), first reported, {cur} millions", "Country codes as tagged by the company (US, CN, JP, ...). Basis as disclosed (ship-to / customer location).")
    rows = [dict(key=n, label=n, values=dict(zip(s.period_end, s["first"])), note=cur + "m") for n, s in a[a.metric == "revenue"].groupby("name")]
    grid(ws, 4, periods, rows, None)

def guidance(t, cfg):
    fpi = cfg.get("filer_type") == "FPI"
    if os.path.exists(f"out/{t}_guid_final.pkl"): return pd.read_pickle(f"out/{t}_guid_final.pkl")
    G = pd.read_pickle(f"out/{t}_guid.pkl") if fpi else pd.read_pickle(f"out/{t}_earn.pkl")["G"]
    G = G.copy(); G["filed"] = pd.to_datetime(G.filed); G["method"] = "regex"
    A = json.load(open("out/model_answers.json"))
    for k, v in A.items():
        task, acc = k.split("|")
        if task != f"guidance_{t}": continue
        o = v["output"]; i = G.index[G.accn == acc]
        if not len(i): continue
        i = i[0]
        if o.get("guide_period") in ("full year", "half year") or o.get("no_guidance"):
            G.loc[i, "method"] = f"model ({v['model']}): {'no quarterly guide' if not o.get('no_guidance') else 'no guidance given'}"; continue
        for c in ("revenue_mid", "revenue_low", "revenue_high", "revenue_qoq_pct", "revenue_direction"):
            if o.get(c) is not None and (c not in G or pd.isna(G.loc[i, c]) if c in G else True): G.loc[i, c] = o[c]
        G.loc[i, "method"] = f"model ({v['model']})" + (" — update of current-quarter guide" if "update" in str(o.get("guide_period")) else "")
    return G

def sheet_guidance(wb, t, cfg, cur):
    from check_guidance import actuals
    ws = wb.create_sheet("Guidance vs actual"); G = guidance(t, cfg); Aq = actuals(t)
    title(ws, f"{t} — company revenue guidance vs what was then reported, {cur} millions",
          "Guided quarter = the quarter after the release (or the current quarter for mid-quarter updates). Actual = first-reported revenue. Implied midpoint, beat and growth are formulas.")
    cols = ["Release filed", "Filing", "Guided quarter end", "Guide low", "Guide high", "Guide midpoint (printed)", "Guide q/q % (printed)", "Direction only",
            "Prior-quarter actual", "Implied midpoint", "Actual revenue", "Actual vs guide %", "Guided q/q growth", "Actual q/q growth", "How read"]
    widths = [12, 22, 12, 11, 11, 12, 11, 12, 12, 12, 12, 11, 11, 11, 34]
    r0 = 4
    for j, (c, w) in enumerate(zip(cols, widths), 1): hdr(ws, r0, j, c, w)
    ws.row_dimensions[r0].height = 42; r = r0 + 1
    for _, g in G.sort_values("filed").iterrows():
        # identical selection to check_guidance: the guided quarter is the first quarter reported after this release, so a
        # mid-quarter update lands on the quarter it updates (that quarter has not been reported yet when the update is issued).
        prev = Aq[Aq.index < g.filed - pd.Timedelta(days=5)]; nxt = Aq[(Aq.index > g.filed) & (Aq.index <= g.filed + pd.Timedelta(days=120))]
        gq = nxt.index[0] if len(nxt) else None
        vals = [g.filed, g.accn, gq, g.get("revenue_low"), g.get("revenue_high"), g.get("revenue_mid"), g.get("revenue_qoq_pct"), g.get("revenue_direction"),
                prev.iloc[-1] if len(prev) else None, None, nxt.iloc[0] if len(nxt) else None]
        has = any(pd.notna(x) for x in vals[3:8] if not isinstance(x, str)) or isinstance(vals[7], str)
        if not has and "model" not in str(g.method): continue
        for j, v in enumerate(vals, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            nf = "yyyy-mm-dd" if j in (1, 3) else ("0.0" if j == 7 else "#,##0;(#,##0);-" if cur == "USD" else "#,##0.0;(#,##0.0);-")
            put(ws, r, j, v, nf, font=F_BL if j in (4, 5, 6, 7, 9, 11) else F_B, fill=FILL_M if "model" in str(g.method) and j in (4, 5, 6, 7, 8) else None)
        ws.cell(r, 10, f'=IF(ISNUMBER(F{r}),F{r},IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),(D{r}+E{r})/2,IF(AND(ISNUMBER(G{r}),ISNUMBER(I{r})),I{r}*(1+G{r}/100),"")))').number_format = "#,##0;(#,##0);-"
        ws.cell(r, 12, f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(J{r}),J{r}<>0),K{r}/J{r}-1,"")').number_format = "0.0%"
        ws.cell(r, 13, f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(I{r}),I{r}<>0),J{r}/I{r}-1,"")').number_format = "0.0%"
        ws.cell(r, 14, f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(I{r}),I{r}<>0),K{r}/I{r}-1,"")').number_format = "0.0%"
        for j in (10, 12, 13, 14): ws.cell(r, j).font = F_B
        put(ws, r, 15, g.method, font=F_I); r += 1
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    put(ws, r + 1, 1, "Actual vs guide far outside ±20% was spot-read against the release: ASML Q4 2017 (guided €2.1B, EUV recognition lifted actual to €2.56B) and Q1 2020 (COVID) are real surprises, not misreads.", font=F_I)
    # other guided lines (FPI: gross margin; annual guides)
    extra = [c for c in G.columns if re.match(r"^(gm_|fy_|h_|eps_|opex|rd|sga)", c)]
    if extra:
        ws2 = wb.create_sheet("Guidance other lines"); title(ws2, f"{t} — other guided lines as printed (gross margin, EPS, opex, annual / half-year guides)")
        d = G[["filed", "accn"] + extra + ["method"]].dropna(how="all", subset=extra)
        table(ws2, d, 4, widths={"accn": 22, "method": 34})

def sheet_customers(wb, t):
    p = f"out/{t}_customers.pkl"
    try: d = pd.read_pickle(p)
    except FileNotFoundError: return
    ws = wb.create_sheet("Customers"); title(ws, f"{t} — customers at or above 10% of revenue (and top-customer aggregates)",
          "From the 10-Q/10-K concentration paragraphs, read by Haiku (Sonnet on validation failure). is_floor = the filing only says 'more than 10%'. 'none >= 10%' rows are explicit statements.")
    d = d.copy(); d["period_end"] = pd.to_datetime(d.period_end)
    d = d[["period_end", "duration", "customer", "type", "pct", "is_floor", "accn", "filed", "model"]].rename(columns={"duration": "duration_months", "pct": "pct_of_revenue", "accn": "first_reported_in"})
    table(ws, d, 4, widths={"customer": 28, "first_reported_in": 22}, nfs={"pct_of_revenue": "0.0"}, fills=[FILL_M] * len(d))

def sheet_fiscal(wb, t):
    x = pd.read_pickle("out/xcheck_fiscal.pkl"); x = x[x.t == t].drop(columns="t")
    ws = wb.create_sheet("Check vs Fiscal.ai"); title(ws, f"{t} — independent cross-check against Fiscal.ai quarterly (2015+)",
          "match = our first-reported or any later printing within 0.5%. Every non-match was traced to a named cause.")
    s = x.status.value_counts().rename_axis("status").reset_index(name="quarters x metrics")
    r = table(ws, s, 4, widths={"status": 70})
    x = x.rename(columns={"fiscal": "Fiscal.ai value", "ours": "our value"})
    table(ws, x, r + 2, widths={"status": 60})
    try:
        sg = pd.read_pickle("out/xcheck_fiscal_seg.pkl"); sg = sg[sg.t == t].drop(columns="t")
        if len(sg):
            ws2 = wb.create_sheet("Check segments vs Fiscal"); title(ws2, f"{t} — segment values vs Fiscal.ai, matched by segment name (2019+)")
            s = sg.status.value_counts().rename_axis("status").reset_index(name="values"); r = table(ws2, s, 4, widths={"status": 45})
            table(ws2, sg, r + 2, widths={"fiscal_name": 40, "status": 40})
    except FileNotFoundError: pass

def sheet_sources(wb, t, cik, accns):
    j, rows = edgar.submissions(cik); f = pd.DataFrame(rows); f = f[f.accessionNumber.isin(accns)].copy()
    f["url"] = [f"https://www.sec.gov/Archives/edgar/data/{cik}/{a.replace('-', '')}/{p}" for a, p in zip(f.accessionNumber, f.primaryDocument)]
    f = f[["filingDate", "form", "reportDate", "accessionNumber", "url"]].sort_values("filingDate")
    ws = wb.create_sheet("Sources"); title(ws, f"{t} — SEC filings used (every value above carries one of these accession numbers)")
    table(ws, f, 4, widths={"url": 95, "accessionNumber": 22})

def summary(wb, t, cfg, cur, lines, checks, gaps):
    ws = wb.active; ws.title = "Summary"
    title(ws, f"{cfg.get('name', t)} ({t}) — SEC-built history, {cur} millions unless noted",
          f"Built {pd.Timestamp.today().date()} from SEC EDGAR (XBRL + filing text). Deterministic rules first; Haiku/Sonnet only where regex cannot read, and only after validation.")
    r = 4; hdr(ws, r, 1, "Tab", 30); hdr(ws, r, 2, "What it holds", 70); hdr(ws, r, 3, "Coverage", 26); r += 1
    for tab, what, cov in lines: put(ws, r, 1, tab, font=F_BB); put(ws, r, 2, what); put(ws, r, 3, cov); r += 1
    r += 1; hdr(ws, r, 1, "Accuracy check", 30); hdr(ws, r, 2, "Result", 70); hdr(ws, r, 3, "Method", 26); r += 1
    for a, b, c in checks: put(ws, r, 1, a, font=F_BB); put(ws, r, 2, b); put(ws, r, 3, c, font=F_I); r += 1
    r += 1; hdr(ws, r, 1, "Known gaps", 30); hdr(ws, r, 2, "Detail", 70); r += 1
    for a, b in gaps: put(ws, r, 1, a, font=F_BB); put(ws, r, 2, b); r += 1
    for rr in ws.iter_rows(min_row=5, max_row=r):
        for c in rr: c.alignment = Alignment(wrap_text=True, vertical="top")

def sheet_txn_kpis(wb):
    e = pd.read_pickle("out/TXN_endmarkets_raw.pkl"); e["period_end"] = pd.to_datetime(e.year.astype(int).astype(str) + "-12-31")
    ws = wb.create_sheet("End markets annual"); title(ws, "TXN — revenue by end market, % (10-K 'markets for our products' chart)",
        "Share of TI revenue (2013+) or of product revenue (2010-12). TI redefined markets in 2013 (Communications/Computing -> Personal electronics, Comms equipment, Enterprise) and added Data center in 2025. Read by Haiku; every % printed and shares sum to ~100.")
    periods = sorted(e.period_end.unique())
    order = e.groupby("market").year.min().sort_values().index
    rows = [dict(key=m, label=m, values=dict(zip(e[e.market == m].period_end, e[e.market == m].pct)), nf="0", flags={p: "model" for p in e[e.market == m].period_end}, note="% of revenue") for m in order]
    rows.append(dict(key="sum", label="Sum of shares (check)", kind="formula", nf="0",
                     formula=lambda col, rm, j, fc: f"=SUM({col}{min(rm[m] for m in order)}:{col}{max(rm[m] for m in order)})"))
    rows.append(dict(kind="section", label="Implied end-market revenue ($M) = share x annual revenue (formulas)"))
    f = pd.read_pickle("out/TXN_fin.pkl"); f["period_end"] = pd.to_datetime(f.period_end); rv = f[(f.canonical == "revenue") & (f.duration_months == 12) & (f.basis == "as_first_reported")]
    rows.append(dict(key="rev", label="Total revenue (10-K)", values=dict(zip(rv.period_end, rv.value)), note="USD m"))
    for m in order:
        rows.append(dict(key=f"imp|{m}", label=f"{m} ($M, implied)", kind="formula", nf="#,##0",
                         formula=(lambda mm: lambda col, rm, j, fc: f'=IF(AND(ISNUMBER({col}{rm[mm]}),ISNUMBER({col}{rm["rev"]})),{col}{rm[mm]}/100*{col}{rm["rev"]},"")')(m)))
    rm, r = grid(ws, 4, periods, rows, None); legend(ws, r + 1)
    c = pd.read_pickle("out/TXN_china.pkl")
    ws2 = wb.create_sheet("China exposure"); title(ws2, "TXN — China share of revenue as stated in the 10-K (regex, clause by clause)",
        "Definitions changed: 2020-21 'shipments to China-based customers'; 2022+ 'end customers headquartered in China'; 2025 adds 'products shipped into China'.")
    table(ws2, c[["year", "measure", "pct", "text", "accn", "filed"]].rename(columns={"pct": "pct_of_revenue"}), 4, widths={"text": 90, "measure": 32, "accn": 22}, nfs={"pct_of_revenue": "0"})

def sheet_asml_kpis(wb):
    q = pd.read_pickle("out/ASML_kpi_q.pkl"); q["period_end"] = pd.to_datetime(q.period_end)
    ws = wb.create_sheet("KPIs quarterly"); title(ws, "ASML — bookings, backlog, systems sold (units) and other operating KPIs, quarterly, first reported",
        "EUR millions unless units. 2011-2014 bookings/backlog were reported EXCLUDING EUV (separate rows). ASML stopped quarterly bookings after Q4 2025 and backlog after 2017. Book-to-bill is a formula.")
    periods = sorted(q.period_end.unique())
    spec = [("bookings_value", "Net bookings", "EUR m"), ("bookings_value_ex_euv", "Net bookings excl. EUV (2011-14 definition)", "EUR m"),
            ("bookings_units", "Net bookings (units)", "units"), ("bookings_units_ex_euv", "Net bookings excl. EUV (units)", "units"),
            ("backlog_value", "Systems backlog", "EUR m"), ("backlog_value_ex_euv", "Systems backlog excl. EUV", "EUR m"), ("backlog_units", "Systems backlog (units)", "units"),
            ("backlog_units_ex_euv", "Systems backlog excl. EUV (units)", "units"), ("backlog_asp", "ASP of systems backlog", "EUR m"),
            ("systems_sold_units", "Lithography systems sold (units)", "units"), ("new_systems_sold_units", "New systems sold (units)", "units"),
            ("used_systems_sold_units", "Used systems sold (units)", "units"), ("asp_system_sales", "ASP of system sales", "EUR m"),
            ("installed_base_mgmt_sales", "Installed Base Management sales", "EUR m"), ("employees_payroll_fte", "Payroll employees (FTE)", "FTE")]
    rows = []
    for k, lab, u in spec:
        s = q[q.metric == k]
        if len(s): rows.append(dict(key=k, label=lab, values=dict(zip(s.period_end, s.value_first)), note=u, nf="#,##0" if u != "EUR m" or k.endswith("asp") else "#,##0", bold=k == "bookings_value"))
    qq = pd.read_pickle("out/ASML_q.pkl"); qq["period_end"] = pd.to_datetime(qq.period_end)
    rows.append(dict(key="sys", label="Net system sales (results release)", values=dict(zip(qq.period_end, qq.system_sales)), note="EUR m"))
    rows += [dict(kind="section", label="Formulas"),
             dict(key="btb", label="Book-to-bill (net bookings / net system sales)", kind="formula", nf="0.00", formula=ratio("bookings_value", "sys")),
             dict(key="byoy", label="Net bookings y/y %", kind="formula", formula=growth("bookings_value", 4))]
    rm, r = grid(ws, 4, periods, rows, 12); legend(ws, r + 1)
    b = pd.read_pickle("out/ASML_breakdown_annual.pkl"); chk = pd.read_pickle("out/ASML_breakdown_check.pkl")
    ws2 = wb.create_sheet("Sales breakdown annual"); title(ws2, "ASML — net sales by technology, end-use, new/used and country, annual (20-F XBRL, first filing that reported each year), EUR millions",
        "Technology/end-use/new-used split NET SYSTEM sales; geography and product/service split TOTAL net sales. Each (breakdown, year) comes from one filing so schemes never mix. Sum checks vs the quarterly releases are on the right of this note in 'Breakdown checks'.")
    periods = sorted(b.period_end.unique()); rows = []
    for kind in ["technology", "end-use", "new/used", "product vs service", "geography"]:
        s = b[b.kind == kind]
        if not len(s): continue
        rows.append(dict(kind="section", label=kind.title()))
        for nm, g in s.groupby("name"):
            rows.append(dict(key=f"{kind}|{nm}", label=nm, values=dict(zip(g.period_end, g["first"])), note="EUR m"))
    tech = [x for x in b[b.kind == "technology"].name.unique()]
    euv = [f"technology|{x}" for x in tech if x.startswith("EUV")]
    rows += [dict(kind="section", label="Formulas"),
             dict(key="euvsh", label="EUV share of technology sales %", kind="formula",
                  formula=lambda col, rm, j, fc: ("=IFERROR((" + "+".join(f"N({col}{rm[k]})" for k in euv if k in rm) + ")/SUM(" + f"{col}{min(rm[f'technology|{x}'] for x in tech)}:{col}{max(rm[f'technology|{x}'] for x in tech)}" + "),\"\")") if euv else None)]
    grid(ws2, 4, periods, rows, None, money="#,##0.0;(#,##0.0);-")
    ws3 = wb.create_sheet("Breakdown checks"); title(ws3, "ASML — annual breakdown sum checks against the four quarterly results releases")
    chk = chk.copy(); chk.loc[(chk.year == 2017) & chk.status.str.startswith("DIFF"), "status"] = "restated basis: first split for 2017 is in the FY2018 20-F (ASC 606 restated); releases are pre-restatement"
    table(ws3, chk, 4, widths={"status": 80, "kind": 20})
    h = pd.read_pickle("out/ASML_h1_raw.pkl"); h["period_end"] = pd.to_datetime(h.period_end)
    ws4 = wb.create_sheet("Sales breakdown H1"); title(ws4, "ASML — first-half breakdown from the statutory interim reports (2011-2022), EUR millions; H2 = full year - H1",
        "Read by Haiku (one table set per release, Sonnet for 2014 where Haiku merged two rows); every block sums to its printed Total. Units = systems. IFRS interim report (same sales as US GAAP).")
    h = h.sort_values(["kind", "name", "period_end", "accn"]).drop_duplicates(["kind", "name", "period_end"])     # first report of each H1
    table(ws4, h[["kind", "period_end", "name", "units", "value", "measure", "accn", "model"]], 4, widths={"name": 28, "accn": 22}, fills=[FILL_M] * len(h))

def build(t):
    """Per-ticker workbook. Every tab is optional: a ticker with one reportable segment gets no segment tabs, an IFRS foreign filer
    gets annual financials only, and anything skipped is listed on the Summary with the reason."""
    cfg = json.load(open(f"config/{t}.json")); cur = cfg.get("currency", "USD"); fpi = cfg.get("filer_type") == "FPI"
    cfg.setdefault("name", t); wb = Workbook(); skipped = []; lines = []; checks = []; gaps = []
    has_q = True
    try:
        qd, f = fin_long(t, cfg)
    except FileNotFoundError:
        raise SystemExit(f"{t}: no financials extracted")
    if not len(qd[qd.duration_months == 3]):
        has_q = False; skipped.append(("Income stmt quarterly", "no quarterly facts in XBRL — IFRS foreign filers tag annual figures only; quarterly needs the 6-K results tables"))
    if has_q:
        sheet_is(wb, t, cfg, cur, qd, "as_first_reported", "Income stmt quarterly", "Values as first reported (the number the market saw). Q4 = FY - 9M where the company files no Q4 report (shaded).")
        if (qd.basis == "latest_filed").any(): sheet_is(wb, t, cfg, cur, qd, "latest_filed", "Income stmt latest", "Latest printing of each quarter (restatements and recasts applied).")
    sheet_annual(wb, t, cfg, cur, f)
    if len(f[f.duration_months == 0]): sheet_bs(wb, t, cfg, cur, f)
    out = {}
    if os.path.exists(f"out/{t}_seg.pkl"):
        s_ = pd.read_pickle(f"out/{t}_seg.pkl")
        if len(s_[(s_["axis"] == "segment") & s_.name.notna()].name.unique()) > 1:
            out = sheet_segments(wb, t, cfg, cur); sheet_geo_annual(wb, t, cfg, cur, out)
        else:
            skipped.append(("Segments", "the company reports a single reportable segment, so there is nothing to break out"))
    if t == "TXN": sheet_txn_kpis(wb)
    if t == "ASML": sheet_asml_kpis(wb)
    try: sheet_guidance(wb, t, cfg, cur)
    except (FileNotFoundError, KeyError) as e: skipped.append(("Guidance vs actual", f"no results releases parsed ({type(e).__name__})"))
    sheet_customers(wb, t)
    if not os.path.exists(f"out/{t}_customers.pkl"): skipped.append(("Customers", "no customer-concentration paragraph found in any filing, or the model read is not in yet"))
    try: sheet_fiscal(wb, t)
    except (FileNotFoundError, KeyError): skipped.append(("Check vs Fiscal.ai", "cross-check not run for this ticker yet"))
    acc = (set(qd.accn) if len(qd) else set()) | set(f.accn)
    for k in ("_seg", "_customers", "_guid", "_kpi_printings"):
        try: acc |= set(pd.read_pickle(f"out/{t}{k}.pkl").accn)
        except Exception: pass
    try: acc |= set(pd.read_pickle(f"out/{t}_earn.pkl")["G"].accn)
    except Exception: pass
    sheet_sources(wb, t, cfg["cik"], acc)
    d3 = qd[(qd.duration_months == 3) & (qd.basis == "as_first_reported") & (qd.canonical == "revenue")] if len(qd) else qd
    ann = f[(f.duration_months == 12) & (f.canonical == "revenue")]
    lines.append(("Income stmt quarterly / latest", "P&L by quarter, first reported and latest printing; margins and growth as formulas",
                  f"{d3.period_end.min().date()} to {d3.period_end.max().date()} ({len(d3)} quarters)" if len(d3) else "not available"))
    lines.append(("Income stmt annual, Balance sheet", f"From XBRL ({cfg.get('annual_form', '10-K')})", f"{ann.period_end.min().date()} to {ann.period_end.max().date()}" if len(ann) else "n/a"))
    if out.get("segment_annual") is not None and "segment" not in out:
        a_ = out["segment_annual"]
        lines.append(("Segments annual", "Reportable segments from the 20-F XBRL (this filer tags no quarterly segment figures)",
                      f"{a_.period_end.min().date()} to {a_.period_end.max().date()}"))
    if out.get("segment") is not None:
        lines.append(("Segments quarterly / detail", "Reportable segments (+ sub-segments), first reported and latest; share and growth formulas",
                      f"{out['segment'].period_end.min().date()} to {out['segment'].period_end.max().date()}"))
        rep = pd.read_pickle(f"out/{t}_addup.pkl") if os.path.exists(f"out/{t}_addup.pkl") else pd.DataFrame()
        if len(rep):
            vc = rep.issue.value_counts(); ok = int(vc.get("sum ok", 0)); bad = int(sum(v for k, v in vc.items() if k.startswith("sum ") and k != "sum ok"))
            checks.append(("Segments add up to total revenue", f"{ok} of {ok + bad} filing x period checks reconcile" + (f"; {bad} still short" if bad else ""), "every filing x period"))
            rules = [f"{k.split('—')[0].strip()}: {v}" for k, v in vc.items() if not k.startswith("sum ")]
            if rules: checks.append(("Rules applied to segment tags", "; ".join(rules), "each is validated by the add-up"))
    if os.path.exists(f"out/{t}_guid_check.pkl"):
        g = pd.read_pickle(f"out/{t}_guid_check.pkl")
        if len(g): checks.append(("Guidance plausibility", "; ".join(f"{k}: {v}" for k, v in g.status.value_counts().items()), "guided-quarter actual within ±20% of implied midpoint"))
    if os.path.exists("out/xcheck_fiscal.pkl"):
        x = pd.read_pickle("out/xcheck_fiscal.pkl"); x = x[x.t == t]
        if len(x): checks.append(("Fiscal.ai cross-check (2015+)", "; ".join(f"{k}: {v}" for k, v in x.status.value_counts().items()), "revenue / gross profit / net income per quarter"))
    if len(cfg.get("cik_chain", [1])) > 1: gaps.append(("Predecessor registrants", cfg.get("_cik_chain_note", "") + " — filings from all of these CIKs are included"))
    if fpi and t != "ASML": gaps.append(("Quarterly detail", "Foreign private issuer: the 20-F carries annual figures, quarterly detail lives in 6-K results releases and needs a per-filer table reader (built for ASML so far)"))
    if not fpi: gaps.append(("Q4", "Q4 = full year minus nine months, from the originally reported figures; a later restatement of the year does not re-derive Q4"))
    gaps.append(("Before 2015", "No Fiscal.ai cross-check available; earlier years rest on the filings agreeing with themselves"))
    for a_, b_ in skipped: gaps.append((f"No '{a_}' tab", b_))
    summary(wb, t, cfg, cur, lines, checks, gaps)
    path = f"out/{t}_sec_history.xlsx"; wb.save(path)
    print(t, "saved", len(wb.worksheets), "tabs", "| skipped:", [x[0] for x in skipped] or "none")
    return path

if __name__ == "__main__":
    for t in sys.argv[1:]:
        try: build(t)
        except SystemExit as e: print(e)
        except Exception as e: print(t, "BUILD ERROR", type(e).__name__, e)
